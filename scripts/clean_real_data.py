"""Run: python clean_real_data.py listings.csv rent.xlsx output_folder"""
import csv
import json
import sys
import unicodedata
from pathlib import Path
from decimal import Decimal
from datetime import date
from collections import Counter, defaultdict
import openpyxl

def text(value):
    return unicodedata.normalize('NFC',str(value)).strip() if value is not None else ''

def integer(value, low=0, high=None):
    if not text(value): return None
    number=Decimal(text(value))
    if number != number.to_integral_value() or number < low or (high is not None and number > high):
        raise ValueError('Invalid integer: '+str(value))
    return int(number)

def boolean(value):
    value=text(value)
    if not value: return None
    if value not in ['t','f']: raise ValueError('Invalid boolean: '+value)
    return int(value=='t')

def money(value):
    if not text(value): return None
    amount=Decimal(text(value).replace('$','').replace(',',''))
    if not amount.is_finite() or amount <= 0 or amount >= 1000000:
        raise ValueError('Invalid price: '+str(value))
    return str(amount.quantize(Decimal('0.01')))

def write(folder,name,fields,records):
    with (folder/(name+'.csv')).open('w',newline='',encoding='utf-8') as stream:
        writer=csv.DictWriter(stream,fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)

def main():
    listing_file,rent_file,output=map(Path,sys.argv[1:])
    output.mkdir(parents=True,exist_ok=True)
    rows=list(csv.DictReader(listing_file.open(encoding='utf-8-sig',newline='')))
    workbook=openpyxl.load_workbook(rent_file,data_only=True)
    aliases={'el Poble Sec':'el Poble Sec - AEI Parc Montjuïc',
             'la Marina del Prat Vermell':'la Marina del Prat Vermell - AEI Zona Franca'}
    neighborhoods=[]
    lookup={}
    for row in workbook['2026'].iter_rows(min_row=22,max_row=94,values_only=True):
        code=integer(row[0],1,73)
        name=text(row[1])
        if code in [n['neighborhood_id'] for n in neighborhoods] or name in lookup:
            raise ValueError('Duplicate neighborhood')
        lookup[name]=code
        neighborhoods.append(dict(neighborhood_id=code,city_id=1,name=name))
    mapping=[]
    for name in sorted(set(text(r['neighbourhood_cleansed']) for r in rows)):
        canonical=aliases.get(name,name)
        if canonical not in lookup: raise ValueError('Unmatched neighborhood: '+name)
        mapping.append(dict(source_name=name,canonical_name=canonical,neighborhood_id=lookup[canonical],rule='explicit alias' if name in aliases else 'NFC and trim'))
    observations=[]
    for year,quarters in [(2025,[1,2,3,4]),(2026,[1])]:
        for row in workbook[str(year)].iter_rows(min_row=22,max_row=94,values_only=True):
            code=integer(row[0],1,73)
            if text(row[1]) not in lookup or lookup[text(row[1])]!=code:
                raise ValueError('Neighborhood identity changed across years')
            for quarter in quarters:
                value=row[quarter+1]
                if value is None: continue
                if not isinstance(value,(int,float)) or value<0: raise ValueError('Invalid rent cell')
                # Zero is not treated as a measured free rent; preserve unknown as NULL.
                observations.append(dict(neighborhood_id=code,observation_date=date(year,3*quarter-2,1).isoformat(),avg_rent_per_m2=str(Decimal(str(value)).quantize(Decimal('0.01'))) if value>0 else None))
    unique={}
    duplicates=0
    for row in rows:
        key=integer(row['id'],1,2**64-1)
        if key in unique:
            if unique[key]!=row: raise ValueError('Conflicting duplicate listing ID')
            duplicates+=1
        else: unique[key]=row
    rows=[unique[key] for key in sorted(unique)]
    host_groups=defaultdict(list)
    for row in rows: host_groups[integer(row['host_id'],1,2**64-1)].append(row)
    hosts=[]
    host_map={}
    conflicts=[]
    for internal,external in enumerate(sorted(host_groups),1):
        host_map[external]=internal
        group=host_groups[external]
        latest=max(date.fromisoformat(r['last_scraped']) for r in group)
        values={boolean(r['host_identity_verified']) for r in group if date.fromisoformat(r['last_scraped'])==latest and text(r['host_identity_verified'])}
        verified=next(iter(values)) if len(values)==1 else None
        if len(values)>1: conflicts.append(external)
        hosts.append(dict(host_id=internal,source_host_id=external,host_type=None,verified=verified,registration_date=None,country_origin=None))
    properties=[]; relationships=[]; listings=[]; snapshots=[]
    rooms={'Entire home/apt':'entire_home','Private room':'private_room','Shared room':'shared_room','Hotel room':'hotel_room'}
    for internal,row in enumerate(rows,1):
        name=aliases.get(text(row['neighbourhood_cleansed']),text(row['neighbourhood_cleansed']))
        host=host_map[integer(row['host_id'],1,2**64-1)]
        properties.append(dict(property_id=internal,neighborhood_id=lookup[name],property_type=text(row['property_type']),bedrooms=integer(row['bedrooms'],0,255),capacity=integer(row['accommodates'],1,255),street_address=None,price_per_m2_purchased=None))
        relationships.append(dict(host_id=host,property_id=internal,ownership_type=None,acquisition_date=None))
        listings.append(dict(listing_id=internal,source_listing_id=integer(row['id'],1,2**64-1),property_id=internal,host_id=host,platform_id=1,room_type=rooms[text(row['room_type'])],license_number=text(row['license']) or None,first_listed_date=None,last_listed_date=None,minimum_nights=integer(row['minimum_nights'],1,65535)))
        snapshots.append(dict(snapshot_id=internal,listing_id=internal,snapshot_date=date.fromisoformat(row['last_scraped']).isoformat(),nightly_price=money(row['price']),available_days_next_365=integer(row['availability_365'],0,365),active=None,reviews_count=integer(row['number_of_reviews'],0,2**32-1)))
    tables={'city':[dict(city_id=1,name='Barcelona',country='Spain')],
            'neighborhood':neighborhoods,'host':hosts,'property':properties,'host_property':relationships,
            'platform':[dict(platform_id=1,name='Airbnb',website_url='https://www.airbnb.com')],
            'listing':listings,'listing_snapshot':snapshots,'housing_market_observation':observations}
    assert len(listings)>=50 and len(observations)>=50
    assert len({(r['neighborhood_id'],r['observation_date']) for r in observations})==len(observations)
    for name,records in tables.items(): write(output,name,list(records[0]),records)
    write(output,'neighborhood_mapping',list(mapping[0]),mapping)
    summary=dict(input_listing_rows=len(unique)+duplicates,unique_listings=len(listings),removed_exact_duplicates=duplicates,
                 rent_observations=len(observations),unique_rent_keys=len(observations),positive_rent_observations=sum(r['avg_rent_per_m2'] is not None for r in observations),zero_rent_cells_as_null=sum(r['avg_rent_per_m2'] is None for r in observations),counts={k:len(v) for k,v in tables.items()},
                 missing_prices=sum(r['nightly_price'] is None for r in snapshots),missing_bedrooms=sum(r['bedrooms'] is None for r in properties),
                 verification_conflict_hosts=len(conflicts),verification_conflict_source_ids=conflicts,all_neighborhood_names_mapped=True,
                 listing_dates=dict(Counter(r['snapshot_date'] for r in snapshots)),currency='EUR',rent_unit='EUR/m2/month')
    (output/'cleaning_summary.json').write_text(json.dumps(summary,indent=2))
    print(json.dumps({k:v for k,v in summary.items() if k!='verification_conflict_source_ids'},indent=2))

if __name__=='__main__': main()
